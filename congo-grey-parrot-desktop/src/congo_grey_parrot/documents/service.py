from __future__ import annotations

from pathlib import Path

from congo_grey_parrot.core.errors import AppError

FONT_CANDIDATES = (
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
    Path("C:/Windows/Fonts/arial.ttf"),
    Path("C:/Windows/Fonts/segoeui.ttf"),
)


def find_pdf_font() -> Path:
    for candidate in FONT_CANDIDATES:
        if candidate.is_file():
            return candidate
    raise AppError(
        "Não encontrei uma fonte Unicode para o PDF. Instale as fontes DejaVu ou Liberation."
    )


def read_document(path: Path) -> str:
    if not path.is_file():
        raise AppError("Arquivo não encontrado.")
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".md", ".srt", ".vtt", ".ass", ".ssa", ".json", ".jsonl", ".tsv"}:
            return path.read_text(encoding="utf-8")
        if suffix == ".docx":
            from docx import Document

            document = Document(str(path))
            return "\n".join(paragraph.text for paragraph in document.paragraphs)
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise AppError("Este PDF está protegido e não pode ser lido aqui.")
            pages = [page.extract_text() or "" for page in reader.pages]
            text = "\n\n".join(part.strip() for part in pages if part.strip())
            if not text:
                raise AppError(
                    "Este PDF não tem texto selecionável. OCR ainda não faz parte deste aplicativo."
                )
            return text
    except AppError:
        raise
    except UnicodeDecodeError as exc:
        raise AppError("O arquivo não está em UTF-8.") from exc
    except Exception as exc:
        raise AppError(f"Não foi possível ler o documento. {exc}") from exc
    raise AppError("Formato não suportado. Use TXT, MD, DOCX, PDF, SRT ou VTT.")


def write_docx(path: Path, text: str, title: str = "") -> None:
    from docx import Document

    document = Document()
    if title:
        document.add_heading(title, level=1)
    for paragraph in text.splitlines():
        document.add_paragraph(paragraph)
    document.save(str(path))


def write_pdf(path: Path, text: str, title: str = "") -> None:
    from fpdf import FPDF

    font = find_pdf_font()
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.add_font("Body", "", str(font))
    pdf.set_font("Body", size=12)
    content = f"{title}\n\n{text}" if title else text
    pdf.multi_cell(0, 7, content)
    pdf.output(str(path))
